# -*- coding: utf-8 -*-
#+++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++
# This is part of the AI engineering lab project
#
# Filename: generateModelComparisonTable.py
#
# Details:  Reads the merged 50-model results from
#           misc/paperTableData/overallResults50models_2026-06-25_mergedLLMs.xlsx
#           and exports two paper-ready LaTeX tables (.tex) into misc/paperOutput:
#             - tableModelComparison.tex    : top-N LLMs, success rates per metric
#             - tableModelComparisonAll.tex : all LLMs (for the appendix)
#
#           Model name convention (mirrors ExuAI.myLLM.ConvertFullName2paperName,
#           the same labels postProcessing.py uses in the PV plots):
#             - proprietary/API models (VRAM < 0): name kept as-is
#               (e.g. Claude-Opus-4.8, GPT-5.2)
#             - open-weight models: ollama tag with registry prefix stripped
#               (e.g. qwen3.6:27b, gpt-oss:20b, qwen3.5-coder:35b)
#           A horizontal rule separates the proprietary group from the
#           open-weight group.
#
# Authors:  Michael Pieber, Claude Opus4.8
# Date:     2026-06-30
# Notes:    Standalone export script. Run from anywhere:
#               python misc/flexuAI/generateModelComparisonTable.py
#
# License:  BSD-3 license
#
# Installation:
#           pip install openpyxl
#
#+++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++

import os

import openpyxl


#%%+++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++
# configuration
#+++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++

scriptDir = os.path.dirname(os.path.abspath(__file__))
miscDir = os.path.dirname(scriptDir)
paperTableDataDir = os.path.join(miscDir, 'paperTableData')
xlsxPath = os.path.join(paperTableDataDir, 'overallResults50models_2026-06-25_mergedLLMs.xlsx')
outputDir = os.path.join(miscDir, 'paperOutput')
outputPath = os.path.join(outputDir, 'tableModelComparison.tex')
outputPathAll = os.path.join(outputDir, 'tableModelComparisonAll.tex')

numberOfModels = 10  # use the first N LLMs (workbook is sorted by overall success)

# table column (header label) -> workbook header
metricColumns = [
    ('comp.', 'choose items'),
    ('solver', 'solver success'),
    ('topol.', 'topology match'),
    ('numsol', 'solution match'),
    ('graph', 'graph content match'),
    ('overall', 'overall success'),
]

# optional footnote markers appended to specific model names (paper name -> marker)
modelFootnotes = {
    'gpt-oss:20b': '\\textsuperscript{$\\dagger$}',
    'gpt-oss:120b': '\\textsuperscript{$\\dagger$}',
}


def ConvertFullNameToPaperName(fullName):
    """
    @brief Strip ollama/registry prefixes from a model full name, mirroring
           ExuAI.myLLM.ConvertFullName2paperName.
    @param fullName workbook 'LLM full name' value
    @return readable model name used in the paper
    """
    for prefix in ('mdq100/', 'meta-llama/', 'Qwen/'):
        if fullName.startswith(prefix):
            fullName = fullName[len(prefix):]
    return fullName


#%%+++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++
# workbook parsing
#+++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++

def ParseModelRows(workbookPath):
    """
    @brief Parse all model data rows of the merged-results workbook.
    @param workbookPath absolute path to the .xlsx file
    @return list of dicts with keys 'name', 'vram' and one per metric label,
            in workbook order (sorted by overall success)
    """
    workbook = openpyxl.load_workbook(workbookPath, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = [list(r) for r in worksheet.iter_rows(values_only=True)]

    # locate the header row (the one containing 'Case Name') and map names -> index
    headerMap = None
    for row in rows:
        if 'Case Name' in row:
            headerMap = {name: idx for idx, name in enumerate(row) if name is not None}
            break
    if headerMap is None:
        raise ValueError("header row with 'Case Name' not found in workbook")

    fullNameIndex = headerMap['LLM full name']
    vramIndex = headerMap['VRAM (GB)']

    models = []
    headerSeen = False
    for row in rows:
        if not headerSeen:
            if 'Case Name' in row:
                headerSeen = True
            continue
        fullName = row[fullNameIndex]
        if not fullName or not str(fullName).strip():
            continue
        record = {
            'name': ConvertFullNameToPaperName(str(fullName)),
            'vram': row[vramIndex],
        }
        for label, header in metricColumns:
            record[label] = row[headerMap[header]]
        models.append(record)
    return models


#%%+++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++
# LaTeX table export
#+++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++

def FormatRate(value, info=None):
    """@brief Format a fraction (0..1) as a percentage with one decimal."""
    if not isinstance(value, (int, float)):
        return '--'
    return f'{100.0 * value:.1f}'


def EscapeLatexText(text):
    """@brief Escape LaTeX-special characters that can occur in model names
       (e.g. the underscores in 'glm-4.7-flash:q4_K_M')."""
    for char, replacement in (('&', '\\&'), ('%', '\\%'), ('$', '\\$'),
                              ('#', '\\#'), ('_', '\\_')):
        text = text.replace(char, replacement)
    return text


def IsProprietary(model):
    """@brief True for API/proprietary models (no local VRAM footprint)."""
    return not isinstance(model['vram'], (int, float)) or model['vram'] < 0


def ComputeGroupMean(models):
    """
    @brief Average each success-rate metric over a group of models.
    @param models list of parsed model records
    @return dict metricLabel -> mean fraction (0..1), or None if no values
    """
    means = {}
    for label, _ in metricColumns:
        values = [m[label] for m in models if isinstance(m[label], (int, float))]
        means[label] = (sum(values) / len(values)) if values else None
    return means


def WriteModelComparisonTable(displayModels, meanRows, outputPath):
    """
    @brief Write the model-comparison table (per-metric success rates), with a
           rule separating proprietary from open-weight models and optional
           group-mean rows appended below the listed models.
    @param displayModels parsed model records to list individually
    @param meanRows list of (label, meansDict) summary rows, or None/empty for none
    @param outputPath absolute path of the .tex file to write
    """
    lines = []
    lines.append('% auto-generated by misc/flexuAI/generateModelComparisonTable.py')
    lines.append('% requires \\usepackage{booktabs} in preamble')
    lines.append('\\begin{tabular}{l@{\\hspace{-5pt}}rrrrrr}')
    lines.append('\\toprule')
    lines.append(' & \\multicolumn{6}{c}{\\textbf{success rate (\\%)}} \\\\')
    lines.append('\\cmidrule(lr){2-7}')
    headerLabels = ' & '.join(f'\\textbf{{{label}}}' for label, _ in metricColumns)
    lines.append(f'\\textbf{{LLM}} & {headerLabels} \\\\')
    lines.append('\\midrule')

    previousProprietary = None
    for model in displayModels:
        # rule between the proprietary group and the open-weight group
        if previousProprietary is True and not IsProprietary(model):
            lines.append('\\midrule')
        previousProprietary = IsProprietary(model)

        name = EscapeLatexText(model['name']) + modelFootnotes.get(model['name'], '')
        rateCells = ' & '.join(FormatRate(model[label]) for label, _ in metricColumns)
        lines.append(f'{name} & {rateCells} \\\\')

    # optional group-mean summary rows
    if meanRows:
        lines.append('\\midrule')
        for label, means in meanRows:
            rateCells = ' & '.join(FormatRate(means[metricLabel]) for metricLabel, _ in metricColumns)
            lines.append(f'{label} & {rateCells} \\\\')

    lines.append('\\bottomrule')
    lines.append('\\end{tabular}')

    os.makedirs(os.path.dirname(outputPath), exist_ok=True)
    with open(outputPath, 'w', encoding='utf-8') as fileHandle:
        fileHandle.write('\n'.join(lines) + '\n')
    print(f'Model-comparison table written to: {outputPath} ({len(displayModels)} models)')


#%%+++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++
# main
#+++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++++

def Main():
    """@brief Parse the workbook and export the model-comparison LaTeX table."""
    if not os.path.isfile(xlsxPath):
        raise FileNotFoundError(f'input workbook not found: {xlsxPath}')
    allModels = ParseModelRows(xlsxPath)
    # sort by overall success (descending) so the listed "first N" are the top
    # performers regardless of the row order in the workbook
    allModels.sort(key=lambda m: m['overall'] if isinstance(m['overall'], (int, float)) else float('-inf'),
                   reverse=True)
    displayModels = allModels[:numberOfModels]
    proprietaryModels = [m for m in allModels if IsProprietary(m)]
    openWeightModels = [m for m in allModels if not IsProprietary(m)]
    topOpenWeight = [m for m in displayModels if not IsProprietary(m)]

    meanRows = [
        ('mean (proprietary)', ComputeGroupMean(proprietaryModels)),
        (f'mean (top {len(topOpenWeight)} open-weight)', ComputeGroupMean(topOpenWeight)),
        (f'mean (all {len(openWeightModels)} open-weight)', ComputeGroupMean(openWeightModels)),
    ]

    print(f'Parsed {len(allModels)} models '
          f'({len(proprietaryModels)} proprietary, {len(openWeightModels)} open-weight); '
          f'listing first {len(displayModels)}.')
    # main table: top-N LLMs with group means
    WriteModelComparisonTable(displayModels, meanRows, outputPath)
    # appendix table: all LLMs, without the mean rows
    WriteModelComparisonTable(allModels, None, outputPathAll)


if __name__ == '__main__':
    Main()
